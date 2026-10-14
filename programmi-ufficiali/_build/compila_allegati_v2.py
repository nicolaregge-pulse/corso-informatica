#!/usr/bin/env python3
"""Allegato A 2026/27, v0.2: inserisce la parte di Regge nei quattro file VERI della scuola (14/10/2026).

Fonte delle voci: programmi-ufficiali/2026-27/PROGRAMMI-DI-MASSIMA-2026-27.md (tabelle dei moduli + sotto-voci).
File di partenza: programmi-ufficiali/2026-27/originali-20261014/ (come li ha mandati Nicola, già toccati da altri docenti).
Regole (Nicola): in ogni area toccata, nella riga dell'annualità di quest'anno, si scrive «Regge» in alto e sotto le nostre
voci; la parte 2026/27 è evidenziata in ROSA; la parte degli anni scorsi resta in verde, gli altri docenti in giallo. Il testo degli altri
docenti non si tocca: la nostra parte va subito sotto l'etichetta dell'anno, prima del resto.
Uso: python3 programmi-ufficiali/_build/compila_allegati_v2.py   (scrive i .docx e stampa il resoconto per l'audit)
"""
import copy, glob, json, os, re, sys
import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

QUI = os.path.dirname(os.path.abspath(__file__))
CART = os.path.join(QUI, '..', '2026-27')
MD = os.path.join(CART, 'PROGRAMMI-DI-MASSIMA-2026-27.md')
COLORE = WD_COLOR_INDEX.PINK  # 14/10 sera, Nicola: parte 2026/27 in rosa; la parte degli anni scorsi resta in verde, gli altri in giallo
# sezione del programma -> (file originale, etichetta dell'anno di quest'anno, nome di uscita)
CLASSI = {
    '03': ('PFP_1', r'^I\s+anno', '1INF'),
    '04': ('PFP_2', r'^II\s+anno', '2INF'),
    '05': ('PFP_3', r'^III\s+anno', '3INF'),
    '06': ('PFP_4', r'^(IV\s+anno|4\s+anno|4°\s+annualit)', '4TI'),
}


def leggi_programma():
    s = open(MD, encoding='utf-8').read()
    prog = {}
    for sec in CLASSI:
        blocco = re.search(r'\n## ' + sec + r' .*?(?=\n## )', s, re.S).group(0)
        moduli = []
        for l in blocco.split('\n'):
            m = re.match(r'^\| (\d+) \| ([^|]+)\| ([^|]+)\| ([^|]+)\|', l)
            if m:
                moduli.append({'n': int(m.group(1)), 'titolo': m.group(2).strip(),
                               'comp': [c.strip() for c in m.group(4).split(',')], 'sotto': []})
        sv = blocco.split('.4 Sotto-voci', 1)[1] if '.4 Sotto-voci' in blocco else ''
        cur = None
        for l in sv.split('\n'):
            m1, m2 = re.match(r'^(\d+)\. (.*)$', l), re.match(r'^   \d+\. (.*)$', l)
            if m1: cur = moduli[int(m1.group(1)) - 1]
            elif m2 and cur: cur['sotto'].append(m2.group(1).strip())
        prog[sec] = moduli
    return prog


def numero_competenza(tab):
    for r in tab.rows[:4]:
        for p in r.cells[0].paragraphs:
            m = re.match(r'^\s*(\d{1,2}(?:\.[a-z])?)\s*[\.\-]\s*\S', p.text)
            if m: return m.group(1)
    return None


def par_dopo(p, testo, grassetto=False):
    nuovo = copy.deepcopy(p._p)
    for figlio in list(nuovo):
        if figlio.tag != qn('w:pPr'): nuovo.remove(figlio)
    p._p.addnext(nuovo)
    q = Paragraph(nuovo, p._parent)
    r = q.add_run(testo); r.bold = grassetto; r.font.highlight_color = COLORE
    return q


def cella_anno(tab, rx):
    """La cella «Argomenti» di quest'anno: la riga che comincia con l'etichetta dell'anno; se manca, la riga vuota dopo «Argomenti»."""
    visti, dopo_arg = set(), False
    for r in tab.rows:
        c = r.cells[0]
        if c._tc in visti: continue
        visti.add(c._tc)
        t = c.text.strip()
        if re.match(rx, t, re.I): return c, True
        if dopo_arg and (t == '' or t.split('\n')[0].strip() == 'Regge'): return c, False  # cella vuota, o già con la nostra parte
        dopo_arg = t == 'Argomenti'
    return None, False


def compila(sec, moduli):
    chiave, rx, nome = CLASSI[sec]
    src = glob.glob(os.path.join(CART, 'originali-20261014', '*%s*' % chiave))[0]
    d = docx.Document(src)
    res = []
    for tab in d.tables[1:]:
        n = numero_competenza(tab)
        if not n: continue
        mie = [m for m in moduli if n in m['comp']]
        if not mie: continue
        cella, ha_etichetta = cella_anno(tab, rx)
        if cella is None:
            res.append({'comp': n, 'esito': 'CELLA NON TROVATA', 'voci': len(mie)}); continue
        altri = '\n'.join(p.text for p in cella.paragraphs[1 if ha_etichetta else 0:]).strip()
        p0 = cella.paragraphs[0]
        if not ha_etichetta and p0.text.strip() == '':
            r = p0.add_run('Regge'); r.bold = True; r.font.highlight_color = COLORE; ultimo = p0
        else:
            ultimo = par_dopo(p0, 'Regge', True)
        for m in mie:
            ultimo = par_dopo(ultimo, m['titolo'])
            for x in m['sotto']:
                ultimo = par_dopo(ultimo, '   – ' + x)
        res.append({'comp': n, 'esito': 'ok' if ha_etichetta else 'cella senza etichetta dell\'anno',
                    'voci': len(mie), 'moduli': [m['n'] for m in mie], 'altri_caratteri': len(altri), 'altri_inizio': altri[:120]})
    out = os.path.join(CART, '20261014_Allegato-A-%s-2026-27_Docente_v0.2.docx' % nome)
    d.save(out)
    usate = sorted({c for m in moduli for c in m['comp']}, key=lambda x: float(x.split('.')[0]))
    messe = {r['comp'] for r in res if r['esito'] != 'CELLA NON TROVATA'}
    return out, res, [c for c in usate if c not in messe]


if __name__ == '__main__':
    prog = leggi_programma()
    tutto = {}
    for sec, moduli in prog.items():
        out, res, mancanti = compila(sec, moduli)
        tutto[CLASSI[sec][2]] = {'file': os.path.basename(out), 'competenze': res, 'competenze_non_trovate': mancanti}
        print(CLASSI[sec][2], os.path.basename(out))
        for r in res: print('   comp %-4s %-30s moduli %s  altri: %s car' % (r['comp'], r['esito'], r.get('moduli'), r.get('altri_caratteri')))
        if mancanti: print('   NON TROVATE:', mancanti)
    json.dump(tutto, open(os.path.join(QUI, 'resoconto_v0.2.json'), 'w'), ensure_ascii=False, indent=1)
